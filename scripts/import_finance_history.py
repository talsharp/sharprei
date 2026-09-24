"""
One-time import of Tal's "Wholesale PNL" Google Sheet export into the Finance
tables. Line items are taken from each monthly tab (so totals are recomputed,
fixing the hand-typed mismatches in the sheet's Total tab), the "Additional"
blocks become General expenses, and the Closed Deals tab becomes the deal list.

Usage:  venv/bin/python scripts/import_finance_history.py "uploads/Wholesale PNL.xlsx" [--apply]
Without --apply it only prints what it would import.
"""

import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.database import SessionLocal, backup_db, init_db  # noqa: E402
from app.finance import get_or_create_vendor  # noqa: E402
from app.models import FinanceExpense, FinanceMonth, Vendor, WholesaleDeal  # noqa: E402

TAB_MONTHS = {
    "August": date(2025, 8, 1),
    "September": date(2025, 9, 1),
    "October": date(2025, 10, 1),
    "November": date(2025, 11, 1),
    "December": date(2025, 12, 1),
    "January": date(2026, 1, 1),
    "February": date(2026, 2, 1),
    "March": date(2026, 3, 1),
    "April": date(2026, 4, 1),
    "May": date(2026, 5, 1),
    "June": date(2026, 6, 1),
    "July": date(2026, 7, 1),
    "August 26": date(2026, 8, 1),
}

# Sheet item name -> (vendor, channel, description prefix)
ITEMS = {
    "chatgpt": ("ChatGPT", "shared", None),
    "highlevel": ("HighLevel", "sms", None),
    "propstream": ("Propstream", "sms", None),
    "launch control": ("Launch Control", "sms", None),
    "batchdata": ("BatchData", "sms", None),
    "batchleads": ("BatchLeads", "sms", None),
    "zapier": ("Zapier", "sms", None),
    "gsuite": ("Gsuite", "shared", None),
    "boots (craig thomas)": ("Boots", "shared", "Craig Thomas"),
    "boots (alex gotiz)": ("Boots", "shared", "Alex Gotiz"),
    "yvonne": ("Yvonne", "sms", None),
    "yvonne (sms + ghl)": ("Yvonne", "sms", "SMS + GHL"),
    "yvonne sms": ("Yvonne", "sms", None),
    "yvonne other": ("Yvonne", "shared", "Other tasks"),
    "dataflik": ("DataFlik", "sms", None),
    "batchdialer": ("BatchDialer", "cold_call", None),
    "joy (cold call)": ("Joy", "cold_call", None),
    "vince (cold call)": ("Vince", "cold_call", None),
    "papers": ("Papers", "shared", None),
    "webwork tracker": ("WebWork Tracker", "sms", None),
    "corey burns": ("Corey Burns", "shared", None),
    "marketing": ("Marketing", "shared", None),
    "incorp services": ("Incorp Services", "shared", None),
    "other": ("Other", "shared", None),
    "facebook": ("Facebook", "facebook", None),
    "hostgator": ("Hostgator", "general", None),
    "microsoft": ("Microsoft", "general", None),
    "credit card cost": ("Credit Card Cost", "general", None),
    "claude": ("Claude", "general", None),
}

RECURRING = {
    "ChatGPT", "HighLevel", "BatchLeads", "WebWork Tracker", "Launch Control", "Zapier", "Gsuite",
    "Vince", "Yvonne", "BatchDialer", "Facebook", "Hostgator", "Microsoft", "Credit Card Cost", "Claude",
}

# (sms normal, sms follow-up, cold call, facebook) - read off each tab's lead boxes
LEADS = {
    date(2025, 9, 1): (1, 0, 0, 0),
    date(2025, 10, 1): (7, 1, 0, 0),
    date(2025, 11, 1): (4, 0, 0, 0),
    date(2025, 12, 1): (0, 4, 0, 0),
    date(2026, 1, 1): (3, 1, 1, 0),
    date(2026, 2, 1): (2, 0, 6, 0),
    date(2026, 3, 1): (9, 1, 4, 0),
    date(2026, 4, 1): (7, 0, 3, 0),
    date(2026, 5, 1): (4, 1, 2, 0),
    date(2026, 6, 1): (2, 0, 1, 0),
    date(2026, 7, 1): (0, 1, 0, 0),
    date(2026, 8, 1): (4, 1, 0, 12),
}

USA_TRIP = (
    date(2026, 9, 1),
    10000,
    "₪28,573 - ESTA, Booking, AirBNB, car rental, flights (Ophir Tours + United, 2 tickets). Excluding food and other.",
)

STATUS = {"YES": "closed", "YES (PORTFOLIO)": "kept", "NO": "cancelled"}
KEPT_ESTIMATES = {"961 Sunglow": 40000, "151 Mahoning": 17500}
SKIP_LABELS = {"item", "total", "leads", "cpl", "normal leads", "follow up leads", "additional"}


def read_month(ws):
    """Yields (item, amount, description, is_additional) for each line item."""
    in_additional = False
    for row in ws.iter_rows(min_row=1, max_col=3, values_only=True):
        name, amount, desc = row
        if not isinstance(name, str):
            continue
        key = name.strip().lower()
        if key == "additional":
            in_additional = True
            continue
        if key in SKIP_LABELS or not isinstance(amount, (int, float)):
            continue
        yield name.strip(), float(amount), (desc.strip() if isinstance(desc, str) else None), in_additional


def parse_sheet_date(value):
    if isinstance(value, datetime):
        return value.date()
    if not isinstance(value, str):
        return None
    cleaned = re.sub(r"(\d+)(st|nd|rd|th)", r"\1", value.strip())
    return datetime.strptime(cleaned, "%d %B %Y").date()


def clean_comment(cell):
    if not cell.comment:
        return None
    text = cell.comment.text
    if "Comment:" in text:
        text = text.split("Comment:", 1)[1]
    return "\n".join(line.strip() for line in text.strip().splitlines() if line.strip())


def num(value):
    return float(value) if isinstance(value, (int, float)) and value else None


def main():
    path = sys.argv[1]
    apply = "--apply" in sys.argv
    wb = openpyxl.load_workbook(path, data_only=True)

    init_db()
    db = SessionLocal()
    if db.query(FinanceMonth).count() or db.query(WholesaleDeal).count():
        sys.exit("Finance data already exists - refusing to import twice.")
    backup = backup_db("before_finance_import") if apply else None

    months = {}
    for tab, month in TAB_MONTHS.items():
        fm = FinanceMonth(month=month)
        normal, fu, cc, fb = LEADS.get(month, (0, 0, 0, 0))
        fm.sms_leads, fm.sms_follow_up_leads, fm.cold_call_leads, fm.facebook_leads = normal, fu, cc, fb
        db.add(fm)
        db.flush()
        months[month] = fm
        wholesale = general = 0.0
        for name, amount, desc, additional in read_month(wb[tab]):
            if amount == 0:
                continue
            mapping = ITEMS.get(name.lower())
            if not mapping:
                sys.exit(f"Unknown item {name!r} in tab {tab!r} - add it to ITEMS first.")
            vendor_name, channel, prefix = mapping
            if additional:
                channel = "general"
            vendor = get_or_create_vendor(db, vendor_name, channel)
            description = " - ".join(p for p in (prefix, desc) if p) or None
            db.add(FinanceExpense(finance_month_id=fm.id, vendor_id=vendor.id, amount=amount, channel=channel, description=description))
            if channel == "general":
                general += amount
            else:
                wholesale += amount
        print(f"{month:%b %Y}: wholesale ${wholesale:,.0f}  general ${general:,.0f}  leads {fm.total_leads}")

    trip_month, trip_amount, trip_desc = USA_TRIP
    fm = FinanceMonth(month=trip_month)
    db.add(fm)
    db.flush()
    vendor = get_or_create_vendor(db, "USA Trip", "general")
    db.add(FinanceExpense(finance_month_id=fm.id, vendor_id=vendor.id, amount=trip_amount, channel="general", description=trip_desc))
    print(f"{trip_month:%b %Y}: general ${trip_amount:,.0f} (USA Trip)")

    for v in db.query(Vendor).all():
        v.recurring = v.name in RECURRING

    ws = wb["Closed Deals"]
    for r in range(2, ws.max_row + 1):
        address_cell = ws.cell(r, 2)
        if not address_cell.value:
            continue
        address = str(address_cell.value).strip()
        closed_raw = ws.cell(r, 5).value
        status = STATUS.get(str(closed_raw).strip().upper(), "under_contract") if closed_raw else "under_contract"
        notes = [clean_comment(address_cell)]
        extra_note = clean_comment(ws.cell(r, 11))
        if extra_note:
            notes.append(f"Additional costs: {extra_note}")
        deal = WholesaleDeal(
            address=address,
            agreement_date=parse_sheet_date(ws.cell(r, 1).value),
            agreement_price=num(ws.cell(r, 3).value),
            sell_price=num(ws.cell(r, 4).value),
            status=status,
            closed_date=parse_sheet_date(ws.cell(r, 6).value),
            revenue=num(ws.cell(r, 7).value) or 0,
            realtor_commission=num(ws.cell(r, 8).value) or 0,
            acquisition_commission=num(ws.cell(r, 9).value) or 0,
            closing_costs=num(ws.cell(r, 10).value) or 0,
            additional_costs=num(ws.cell(r, 11).value) or 0,
            notes="\n\n".join(n for n in notes if n) or None,
        )
        if status == "kept":
            deal.estimated_value = next((v for k, v in KEPT_ESTIMATES.items() if address.startswith(k)), None)
        db.add(deal)
        print(f"Deal: {address} [{status}] profit={deal.profit} est={deal.estimated_value}")

    if apply:
        db.commit()
        print(f"\nImported. Backup of the database from before the import: {backup.name}")
    else:
        db.rollback()
        print("\nDry run only - nothing saved. Re-run with --apply to import.")
    db.close()


if __name__ == "__main__":
    main()
